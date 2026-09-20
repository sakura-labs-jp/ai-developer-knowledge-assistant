from pathlib import Path

from openpyxl import load_workbook

from backend.loaders.base import DocumentRecord


def load_excel(
    file_path: str,
    sheet_name: str | None = None,
) -> list[DocumentRecord]:
    """
    Load rows from an Excel worksheet and convert them
    into normalized DocumentRecord objects.

    The first row is treated as the header.
    Empty rows are ignored.

    This layer performs extraction only.
    It does not sanitize or generate Knowledge.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {file_path}"
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = (
            workbook[sheet_name]
            if sheet_name
            else workbook.active
        )

        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            return []

        headers = [
            str(value).strip()
            if value is not None
            else f"column_{index}"
            for index, value in enumerate(
                header_row,
                start=1,
            )
        ]

        records = []

        for row_index, row in enumerate(
            rows,
            start=2,
        ):
            if all(
                value is None
                for value in row
            ):
                continue

            data = {
                header: value
                for header, value in zip(
                    headers,
                    row,
                )
            }

            records.append(
                DocumentRecord(
                    source_type="Excel",
                    source_location=str(path),
                    source_title=path.name,
                    record_index=row_index,
                    data=data,
                )
            )

        return records

    finally:
        workbook.close()